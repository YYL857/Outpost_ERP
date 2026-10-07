# -*- coding: utf-8 -*-
"""In-place update of existing architecture Word doc (functional sections only)."""
from docx import Document

PATH = r"C:\Users\cjl\Desktop\Outpost-ERP架构与实现方案.docx"


def set_para_text(p, text: str) -> None:
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def main() -> None:
    doc = Document(PATH)

    for p in doc.paragraphs:
        t = p.text

        if t.startswith("「慧运营」是一套面向"):
            set_para_text(
                p,
                "「慧运营」是一套面向 NOON（中东电商平台）卖家的跨境电商经营后台。"
                "系统自己不创造订单和库存，全部平台经营数据以 NOON 为唯一数据源；"
                "本地负责采集归类、同步入库、成本配置、经营可视化和供应链协同。"
                "关键约束：NOON 官方通常不提供可用的开放接口，但可通过网页登录卖家后台。"
                "因此主数据入口为门户模式（portal）：从后台导出 CSV/JSON 导入，"
                "或用登录 Cookie + 数据 URL 拉取页面/XHR 数据，清洗归类后写入 Inbox 再同步到业务库；"
                "业务模块只读本地已同步数据。另保留 mock（本地模拟）与 live（仅当确有官方 API 时）备用。",
            )
        elif t == "在什么环境运行、如何启动与上线对接":
            set_para_text(p, "在什么环境运行、如何启动；没有官方 API 时数据如何采集入库")
        elif t == "可扩展：业务页不直连 NOON；换真实 API 只动接入层。":
            set_para_text(p, "可扩展：业务页不直连 NOON；换采集方式（portal/live/mock）只动接入层。")
        elif t == "可运维：环境变量切换 mock/live；上线清单文档化。":
            set_para_text(p, "可运维：环境变量切换 portal/mock/live；上线清单文档化。")
        elif t.startswith("本系统的平台经营数据一律来自 NOON"):
            set_para_text(
                p,
                "本系统的平台经营数据一律来自 NOON。"
                "在没有官方开放接口的前提下，数据不是接口推过来，而是从卖家网页后台拿到再归类同步。"
                "本地配置（运费、测评费率、上架额度）属于商家自己的经营参数。"
                "边界必须清楚：平台数据只能走接入层进入 Inbox，不能业务页各自抓网页。",
            )
        elif t == "2.2 NOON 进入系统的两条路":
            set_para_text(p, "2.2 没有官方 API 时：门户采集怎么进系统")
        elif t.startswith("推（Webhook）"):
            set_para_text(
                p,
                "主路径——导出文件导入：在 NOON 卖家后台导出订单/商品/库存 CSV 或 JSON，"
                "到慧运营「NOON 接入」页上传或粘贴，系统自动解析归类，写入 Inbox，再幂等同步到业务库。",
            )
        elif t.startswith("拉（增量同步）"):
            set_para_text(
                p,
                "可选路径——Cookie 会话拉取：登录后台后复制 Cookie，并把列表数据的 XHR/页面 URL 配进接入页，"
                "一键拉取后同样进入 Inbox。浏览器自动化（Playwright）代码位已预留，按需启用。",
            )
        elif t.startswith("两条路最后都进同一个 Inbox"):
            set_para_text(
                p,
                "不论导出导入还是会话拉取，最后都进同一个 Inbox，用同一个幂等键去重，"
                "因此重复导入同一批数据不会重复记账。",
            )
        elif t == "2.3 当前开发态 vs 上线态":
            set_para_text(p, "2.3 三种运行模式")
        elif t.startswith("切换方式：复制 .env.example"):
            set_para_text(
                p,
                "切换方式：在「NOON 接入」页选择 portal/mock/live，"
                "或设置环境变量 NOON_MODE=portal|mock|live|inherit。"
                "日常使用优先 portal。详见项目内 GO_LIVE.md。",
            )
        elif t == "平台接入层（platforms/noon）：唯一允许对话 NOON 的地方。":
            set_para_text(
                p,
                "平台接入层（platforms/noon）：唯一允许接触 NOON 的地方"
                "（含 portal 解析/导入/会话拉取，以及 mock/live）。",
            )
        elif t.startswith("NOON →（Webhook 或 增量拉取）"):
            set_para_text(
                p,
                "NOON 卖家后台 →（导出文件导入 或 Cookie+URL 拉取）→ 解析归类为统一 DTO → "
                "Inbox 原始事件落库 → 幂等 UPSERT 写入订单/商品/库存 → "
                "Catalog/销量/财务等业务模块只读本地库 → 同步健康中心展示任务与告警。",
            )
        elif t.startswith("口诀：业务模块永不直连 NOON"):
            set_para_text(
                p,
                "口诀：业务模块永不直连网页；所有外部取数只走接入层；"
                "是否成功以 Inbox 处理结果和健康中心为准，不以文件上传成功本身为准。",
            )
        elif t == "4.2 NOON 接入实现":
            set_para_text(p, "4.2 NOON 接入实现（门户为主）")
        elif t.startswith("接口契约（NoonClient）约定"):
            set_para_text(
                p,
                "接入契约（NoonClient）统一约定：校验就绪、获取订单/商品/库存。"
                "portal 实现优先使用门户缓冲（导出导入或会话拉取写入的数据），再进入 Inbox 管道。",
            )
        elif t == "MockNoonClient：返回稳定可演示的假数据，支撑业务联调。":
            set_para_text(
                p,
                "Portal 导入：后台导出 CSV/JSON →「NOON 接入」上传 → 自动归类同步（推荐主路径）。",
            )
        elif t.startswith("LiveNoonClient："):
            set_para_text(
                p,
                "Portal 会话拉取：配置 Cookie + 订单/商品/库存数据 URL，一键拉取并同步。",
            )
        elif t.startswith("Webhook：POST /api/noon/webhook，校验密钥"):
            set_para_text(
                p,
                "Mock/Live 备用：mock 本地演示；live 仅在确有官方 API 时启用。",
            )
        elif t.startswith("轮询：进程内定时拉取"):
            set_para_text(
                p,
                "浏览器自动化占位：需要全自动抓页时可再启用 Playwright 采集器。",
            )
        elif t == "Catalog：读本地 catalog + stock，可一键从 NOON 同步。":
            set_para_text(
                p,
                "Catalog：读本地 catalog + stock；先在接入页导入/拉取，再刷新查看。",
            )
        elif t == "Webhook：POST /api/noon/webhook":
            set_para_text(
                p,
                "门户导入：前端「NOON 接入」页，或 POST /api/noon/portal/import",
            )
        elif t == "Webhook 必须公网可达；密钥与平台配置一致。":
            set_para_text(
                p,
                "门户 Cookie 属于敏感信息，仅保存在服务端配置，勿提交代码仓库。",
            )
        elif t == "6.2 上线当天最小动作清单":
            set_para_text(p, "6.2 日常使用最小动作清单（门户模式）")
        elif t == "复制并填写 .env，NOON_MODE=live":
            set_para_text(p, "设置 NOON_MODE=portal（或界面选择 portal）并启动服务")
        elif t == "按官方文档核对 endpoints.ts 与鉴权头":
            set_para_text(p, "打开「NOON 接入」，下载模板并对照后台导出字段")
        elif t == "启动服务，打开「NOON 接入」点校验":
            set_para_text(p, "从 NOON 卖家后台导出订单/商品/库存并导入同步")
        elif t == "配置平台 Webhook 到 /api/noon/webhook":
            set_para_text(p, "（可选）配置 Cookie + 数据 URL，使用会话拉取")
        elif t == "在同步健康中心分别同步 orders / catalog / stock":
            set_para_text(p, "在同步健康中心确认 Inbox 处理成功、无红灯")
        elif t == "检查 Catalog、销量、财务页；告警区无红灯":
            set_para_text(p, "检查 Catalog、销量、财务页是否出现对应数据")
        elif t == "高：真实 Live 联调通过 + 对账任务固化":
            set_para_text(
                p,
                "高：按真实后台导出样例固化字段映射；需要时再启用浏览器自动采集",
            )
        elif t.startswith("慧运营 的定位很明确"):
            set_para_text(
                p,
                "慧运营 的定位很明确：它是 NOON 卖家的本地经营驾驶舱。"
                "在没有官方开放接口时，通过门户采集（导出导入 / 会话拉取）把后台数据稳定搬进来，"
                "再叠加上商家自己的成本规则与经营动作。"
                "架构上把「接触 NOON」收口到接入层，把「不能丢、不能错、出错必须看见」做成同步底座默认能力；"
                "网页可以变、导出格式可以调，换的是采集适配器，不是整套 ERP。",
            )

    # tables
    t0 = doc.tables[0]
    t0.rows[1].cells[1].text = "后台里有的单子/库存，同步后系统里最终也要有"
    t0.rows[1].cells[2].text = "先写 Inbox 再处理；导出导入/会话拉取失败可重试重放"
    t0.rows[2].cells[2].text = "幂等写入（订单号/SKU 唯一）；带更新时间；解析表头别名兼容"
    t0.rows[3].cells[2].text = "同步健康中心红黄绿灯；导入/拉取失败进告警列表"

    for row in doc.tables[1].rows:
        if row.cells[0].text.strip() == "NOON 接入":
            row.cells[2].text = (
                "portal/mock/live；导出文件导入；Cookie 会话拉取；模板下载"
            )
            break

    doc.tables[2].rows[1].cells[2].text = "NOON 卖家后台（导出文件 / 网页会话采集）"

    t3 = doc.tables[3]
    while len(t3.rows) < 4:
        t3.add_row()
    t3.rows[1].cells[0].text = "portal（推荐）"
    t3.rows[1].cells[1].text = "官方无开放接口（当前真实情况）"
    t3.rows[1].cells[2].text = "导出导入 / Cookie 会话拉取 → 归类同步"
    t3.rows[2].cells[0].text = "mock"
    t3.rows[2].cells[1].text = "本地演示、暂无后台导出文件时"
    t3.rows[2].cells[2].text = "返回模拟订单/商品/库存，全链路可测"
    t3.rows[3].cells[0].text = "live"
    t3.rows[3].cells[1].text = "仅当将来确有官方 Partner API"
    t3.rows[3].cells[2].text = "走 HTTP Client；路径与鉴权按官方文档配置"

    for row in doc.tables[4].rows:
        cell0 = row.cells[0].text
        if "platforms/noon" in cell0:
            row.cells[0].text = "apps/api/src/platforms/noon/"
            row.cells[1].text = "mock/live Client；portal 采集（parsers/import/session）"
            break

    for row in doc.tables[5].rows:
        if row.cells[0].text.strip() == "平台对接":
            row.cells[1].text = "NoonClient + portal/mock/live"
            row.cells[2].text = "无官方 API 也能用；有 API 可插拔"
            break

    risks = [
        ("重复导入同一文件", "eventId 去重，业务 UPSERT", "Inbox 标记 duplicate / 不重复记账"),
        ("CSV 表头不匹配", "表头别名映射；解析 0 行则报错", "导入失败红灯 + 明确文案"),
        ("Cookie 过期", "会话拉取失败告警", "红灯：需重新登录复制 Cookie"),
        ("后台页面改版", "改 URL/解析器，不动业务页", "健康中心持续报错直到修好"),
        ("处理到一半中断", "事件已在 Inbox，可继续处理", "积压数上升，黄灯"),
    ]
    t6 = doc.tables[6]
    for idx, (a, b, c) in enumerate(risks, start=1):
        if idx < len(t6.rows):
            t6.rows[idx].cells[0].text = a
            t6.rows[idx].cells[1].text = b
            t6.rows[idx].cells[2].text = c

    t8 = doc.tables[8]
    t8.rows[1].cells[1].text = "portal / mock / live / inherit"
    t8.rows[1].cells[2].text = "portal"
    for row in t8.rows:
        if "API_KEY" in row.cells[0].text:
            row.cells[1].text = "仅 live 模式需要的平台密钥"
            row.cells[2].text = "有官方 API 时再填"
        if "WEBHOOK_SECRET" in row.cells[0].text:
            row.cells[1].text = "可选回调密钥（非门户主路径）"
            row.cells[2].text = "一般 portal 可不填"
        if row.cells[0].text.strip() == "NOON_BASE_URL":
            row.cells[1].text = "后台相关根地址（可选）"
            row.cells[2].text = "卖家后台域名"

    doc.tables[9].rows[1].cells[2].text = "主页、配置、NOON 接入（portal）、同步健康、Inbox"

    for row in doc.tables[10].rows:
        key = row.cells[0].text.strip()
        if key == "不丢数":
            row.cells[1].text = "导入已知行数的导出文件后对比本地条数"
            row.cells[2].text = "业务表条数对齐；失败可重导"
        elif key == "不错数":
            row.cells[1].text = "重复导入同一文件；乱序旧数据"
            row.cells[2].text = "不重复记账；旧包不覆盖新状态"
        elif key == "错误可见":
            row.cells[1].text = "上传错误格式文件或过期 Cookie"
            row.cells[2].text = "健康中心/接入页出现明确警告"
        elif key == "即插即用":
            row.cells[0].text = "门户可用"
            row.cells[1].text = "仅用导出导入（不依赖官方 API）"
            row.cells[2].text = "业务页能显示同步后的数据"

    # ensure version untouched
    version_ok = any("版本：v0.1" in p.text for p in doc.paragraphs)
    if not version_ok:
        raise RuntimeError("version line missing or changed unexpectedly")

    import os
    import shutil
    import tempfile

    tmp_dir = tempfile.gettempdir()
    tmp_path = os.path.join(tmp_dir, "Outpost-ERP-architecture-patched.docx")
    doc.save(tmp_path)

    try:
        shutil.copyfile(tmp_path, PATH)
        print("updated:", PATH)
    except PermissionError:
        # Word 占用时写到同目录临时名，提醒关闭后替换
        desktop = os.path.dirname(PATH)
        fallback = os.path.join(desktop, "Outpost-ERP架构与实现方案.待替换.docx")
        shutil.copyfile(tmp_path, fallback)
        print("LOCKED:", PATH)
        print("fallback:", fallback)
        raise SystemExit(2)

    print("version_ok:", version_ok)


if __name__ == "__main__":
    main()
