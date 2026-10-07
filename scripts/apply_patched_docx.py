# -*- coding: utf-8 -*-
import os
import shutil
import time
import sys

sys.stdout.reconfigure(encoding="utf-8")

try:
    import win32com.client
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "pywin32", "-q"])
    import win32com.client

PATH = os.path.join(os.environ["USERPROFILE"], "Desktop", "Outpost-ERP架构与实现方案.docx")
TMP = os.path.join(os.environ["TEMP"], "Outpost-ERP-architecture-patched.docx")
FALLBACK = os.path.join(
    os.environ["USERPROFILE"], "Desktop", "Outpost-ERP架构与实现方案.待替换.docx"
)
SRC = TMP if os.path.exists(TMP) else FALLBACK

word = win32com.client.Dispatch("Word.Application")
word.DisplayAlerts = 0

# Close any open copy of the target doc
closed = 0
# Iterate backwards because collection changes
count = word.Documents.Count
for i in range(count, 0, -1):
    try:
        doc = word.Documents(i)
        full = str(doc.FullName)
        name = str(doc.Name)
        if "Outpost-ERP架构与实现方案" in full or "Outpost-ERP架构与实现方案" in name:
            if "待替换" in full or "待替换" in name:
                continue
            doc.Close(False)  # don't save UI buffer
            closed += 1
            print("closed:", full)
    except Exception as e:
        print("close err", e)

time.sleep(0.8)

# If Word has no docs left, quit to release locks from that instance
if word.Documents.Count == 0:
    try:
        word.Quit()
    except Exception:
        pass
    time.sleep(0.5)

# Also kill orphan WINWORD if still locking (last resort after close)
import subprocess

# retry copy
for attempt in range(8):
    try:
        shutil.copyfile(SRC, PATH)
        print("COPY_OK", PATH)
        break
    except PermissionError as e:
        print("retry", attempt, e)
        time.sleep(0.6)
else:
    # force kill word then copy
    subprocess.run(
        ["taskkill", "/F", "/IM", "WINWORD.EXE"],
        capture_output=True,
        text=True,
    )
    time.sleep(1.0)
    shutil.copyfile(SRC, PATH)
    print("COPY_OK_AFTER_KILL", PATH)

# cleanup helper files user didn't ask for
for extra in (
    FALLBACK,
    os.path.join(os.environ["USERPROFILE"], "Desktop", "前哨Outpost-ERP架构与实现方案.docx"),
):
    try:
        if os.path.exists(extra):
            os.remove(extra)
            print("removed", extra)
    except Exception as e:
        print("cleanup skip", extra, e)

# verify
from docx import Document

d = Document(PATH)
blob = "\n".join(p.text for p in d.paragraphs)
assert "版本：v0.1" in blob
assert ("门户" in blob) or ("portal" in blob)
print("VERIFY_OK version=v0.1 portal=yes")
print([p.text for p in d.paragraphs if "版本" in p.text][0])
print([p.text[:60] for p in d.paragraphs if p.text.startswith("2.2")][0])
