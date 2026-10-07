# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

out_path = os.path.join(
    os.environ["USERPROFILE"], "Desktop", "慧运营Outpost-ERP架构与实现方案.docx"
)

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)


def set_run_font(run, size=11, bold=False, color=None, name="微软雅黑"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_cn(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(
            run,
            size={1: 18, 2: 14, 3: 12}.get(level, 11),
            bold=True,
            name="微软雅黑",
        )
    return p


def add_para(text, size=11, bold=False, space_after=8, first_line=False):
    p = doc.add_paragraph()
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_bullet(text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    for run in p.runs:
        set_run_font(run, size=11)
    return p


def shade_cell(cell, hex_color="D7EBE6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                set_run_font(run, size=10, bold=True)
        shade_cell(hdr[i], "D7EBE6")
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = str(val)
            for p in cells[c_idx].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()
    return table


# Cover
for _ in range(3):
    doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("慧运营")
set_run_font(r, size=28, bold=True, color=(13, 110, 99))

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("跨境电商 ERP 后台系统")
set_run_font(r, size=18, bold=True)

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub2.add_run("架构设计与实现方案说明")
set_run_font(r, size=16, color=(60, 80, 75))

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_before = Pt(24)
r = meta.add_run(
    "项目代号：慧运营 ERP\n"
    "文档用途：架构说明 · 实现方案 · 上线对接指引\n"
    "版本：v0.2（门户采集为主 · 无官方开放接口场景）"
)
set_run_font(r, size=11, color=(90, 100, 98))

doc.add_page_break()

add_heading_cn("写在前面：一句话看懂这个项目", 1)
add_para(
    "「慧运营」是一套面向 NOON（中东电商平台）卖家的跨境电商经营后台。"
    "系统自己不创造订单和库存，全部平台经营数据以 NOON 为唯一数据源；"
    "本地负责采集归类、同步入库、成本配置、经营可视化和供应链协同。",
    first_line=True,
)
add_para(
    "关键现实约束：NOON 官方通常不提供可用的开放接口，但卖家可以通过网页登录后台。"
    "因此本系统的主数据入口是「门户模式（portal）」——"
    "从卖家后台导出文件导入，或用登录会话（Cookie + 数据 URL）拉取页面/XHR 数据，"
    "再清洗归类后同步到慧运营。业务模块只读本地已同步数据，不直连网页。",
    first_line=True,
)
add_para("本文档重点回答：", bold=True)
add_bullet("要解决什么问题、必须达到哪些硬性需求")
add_bullet("没有官方 API 时，数据怎么进系统")
add_bullet("系统怎么分层、模块如何划分、数据怎么流动")
add_bullet("在什么环境运行、如何启动与上线使用")
add_bullet("分几期做、现在做到哪、还要补什么")

add_heading_cn("一、项目目标与核心需求", 1)
add_heading_cn("1.1 业务目标", 2)
add_para(
    "搭建一个简约清晰、一眼能看懂的跨境 ERP 后台，减少平台后台与表格之间来回切换。"
    "卖家可以在一个界面里完成：看店铺状态、管成本配置、把 NOON 后台数据同步进来、"
    "做选品/上架、看销量与利润、推进采购补货发货与活动测评。",
    first_line=True,
)

add_heading_cn("1.2 三条硬性需求（同步质量）", 2)
add_para("这是整个技术架构的出发点和验收标准：", first_line=True)
add_table(
    ["硬性要求", "通俗理解", "工程怎么保证"],
    [
        [
            "不能丢数据",
            "后台里有的单子/库存，同步后系统里最终也要有",
            "先写 Inbox 原始事件再处理；导入/拉取失败可重试重放",
        ],
        [
            "同步不能出错",
            "不能用旧数据覆盖新数据，归类字段要正确",
            "幂等写入（订单号/SKU 唯一）；带更新时间，旧包丢弃；解析器做表头别名兼容",
        ],
        [
            "出错必须看得见",
            "采集失败、解析失败不能悄悄过去",
            "同步健康中心红黄绿灯；告警列表；支持确认与重试",
        ],
    ],
)

add_heading_cn("1.3 功能需求清单", 2)
add_para(
    "界面风格：自有品牌「慧运营」，简约简洁，信息层级清楚；"
    "参考图仅作业务模块参考，不照搬外观与名称。",
    first_line=True,
)
add_table(
    ["模块", "阶段", "要做什么", "当前状态"],
    [
        ["我的主页", "P0", "店铺信息、上架额度、运费/测评配置、同步摘要", "已实现"],
        ["NOON 接入", "P0", "portal/mock/live；导出导入；Cookie 会话拉取", "已实现"],
        ["同步健康中心", "P0", "任务、Inbox、告警、触发同步", "已实现"],
        ["配置管理", "P0", "空运/海运/汇率/体积重；测评佣金/汇率/快递费/余额", "已实现"],
        ["Catalog 产品库", "P1", "展示已同步商品与库存", "已实现"],
        ["链接上架", "P1", "上架任务 + 本月额度扣减", "已实现骨架"],
        ["销量统计", "P1", "订单汇总 GMV/件数/状态", "已实现"],
        ["Best Sell", "P1", "爆款粗排", "已实现骨架"],
        ["财务报表", "P2", "收入来自订单，成本吃配置估算利润", "已实现骨架"],
        ["采购/补货/发货", "P2", "供应链闭环", "路由占位"],
        ["广告/测评活动", "P2", "广告与测评费用", "路由占位"],
        ["选品/多渠道上架", "P3", "选品助手、亚马逊/1688 等", "路由占位"],
    ],
)

add_heading_cn("1.4 非功能需求", 2)
add_bullet("可用性：主流程少步骤，状态用颜色直接表达（绿正常 / 黄延迟 / 红失败）。")
add_bullet("可扩展：业务页不直连 NOON；换采集方式只动接入层（portal/live/mock）。")
add_bullet("可运维：环境变量切换模式；上线清单文档化。")
add_bullet("可替换存储：默认本地 JSON；预留 PostgreSQL 仓储接口。")
add_bullet("页面改版容忍：导出模板 + 表头别名；选择器/URL 可配置，避免写死难改。")

add_heading_cn("二、数据来源说明（最重要）", 1)
add_para(
    "平台经营数据一律来自 NOON 卖家侧。"
    "在没有官方开放接口的前提下，数据不是“接口推过来”，而是“我们从网页后台拿到再归类同步”。"
    "本地配置（运费、测评费率、上架额度）属于商家自己的经营参数，不是平台数据。",
    first_line=True,
)

add_heading_cn("2.1 两类数据", 2)
add_table(
    ["类型", "例子", "来源", "存放位置"],
    [
        [
            "平台数据",
            "订单、Catalog、库存、履约状态",
            "NOON 卖家后台（导出/网页会话采集）",
            "本地业务表（经 Inbox）",
        ],
        [
            "经营配置",
            "运费、汇率、测评佣金、上架额度",
            "商家在慧运营后台配置",
            "本地配置表",
        ],
    ],
)

add_heading_cn("2.2 没有官方 API 时：门户采集三条路", 2)
add_para("系统已实现的主路径如下（推荐优先级从上到下）：", first_line=True)
add_table(
    ["方式", "怎么做", "稳定性", "适用场景"],
    [
        [
            "导出文件导入（推荐）",
            "后台导出 CSV/JSON → 慧运营「NOON 接入」上传 → 自动解析归类 → Inbox → 业务库",
            "最高",
            "日常同步、可人工操作、字段相对固定",
        ],
        [
            "Cookie 会话拉取",
            "登录后台复制 Cookie，填写列表数据的 XHR/页面 URL → 一键拉取并同步",
            "中",
            "能在开发者工具里抓到数据请求地址时",
        ],
        [
            "浏览器自动化（占位）",
            "Playwright 登录并抓页（代码位已留，按需启用）",
            "较低（怕改版）",
            "无法导出、又必须全自动时",
        ],
    ],
)
add_para(
    "无论哪条路，最后都进入同一个 Inbox，用同一个幂等键去重，"
    "因此不会因为多次导入同一批数据而重复记账（相同事件标识会被识别）。",
    first_line=True,
)

add_heading_cn("2.3 三种运行模式", 2)
add_table(
    ["模式", "何时用", "行为"],
    [
        [
            "portal（推荐）",
            "官方无开放接口（当前真实情况）",
            "导出导入 / Cookie 会话拉取 → 归类同步",
        ],
        ["mock", "本地演示、无任何后台文件时", "返回模拟订单/商品/库存，全链路可测"],
        [
            "live",
            "仅当将来确有官方 Partner API",
            "走 HTTP Client；路径与鉴权按官方文档配置",
        ],
    ],
)
add_para(
    "切换方式：在「NOON 接入」页选择模式，或在环境变量中设置 NOON_MODE=portal|mock|live|inherit。"
    "详见项目内 GO_LIVE.md。",
    first_line=True,
)

add_heading_cn("2.4 导出字段建议（导入模板）", 2)
add_bullet("订单：order_nr, status, currency, total_amount, item_count, placed_at, updated_at")
add_bullet("商品：partner_sku, title, is_active, updated_at")
add_bullet("库存：partner_sku, warehouse_code, qty, updated_at")
add_para(
    "表头支持常见别名（如 sku、order id、quantity 等）。"
    "可在「NOON 接入」页下载对应 CSV 模板对照后台导出字段。",
    first_line=True,
)

add_heading_cn("三、总体架构（一看就懂）", 1)
add_heading_cn("3.1 分层结构", 2)
add_para("可以把它想成四层楼：", first_line=True)
add_bullet("展示层（apps/web）：React 简约后台，侧栏模块 + 多标签页。")
add_bullet("接口层（apps/api）：Express 提供 REST API，按业务模块拆路由。")
add_bullet(
    "平台接入层（platforms/noon）：唯一允许接触 NOON 的地方——"
    "含 portal 解析/导入/会话拉取，以及可选 live API、mock。"
)
add_bullet("数据层：Inbox 原始事件 + 业务表；当前 JSON，可换 PostgreSQL。")

add_heading_cn("3.2 核心数据流", 2)
add_para(
    "NOON 卖家后台 →（导出文件 或 Cookie+URL 拉取）→ 解析归类为统一 DTO → "
    "Inbox 原始事件落库 → 幂等 UPSERT 写入订单/商品/库存 → "
    "Catalog/销量/财务等业务模块只读本地库 → 同步健康中心展示任务与告警。",
    first_line=True,
)
add_para(
    "口诀：业务模块永不直连网页；所有外部取数只走接入层；"
    "是否成功以 Inbox 处理结果和健康中心为准，不以“文件上传成功”本身为准。",
    bold=True,
)

add_heading_cn("3.3 工程目录", 2)
add_table(
    ["路径", "作用"],
    [
        ["项目根目录/", "pnpm monorepo 根目录"],
        ["apps/web/", "前端：Vite + React + TypeScript"],
        ["apps/api/", "后端：Express + TypeScript"],
        ["apps/api/src/platforms/noon/portal/", "门户采集：解析、导入、会话拉取、浏览器占位"],
        ["apps/api/src/platforms/noon/", "mock/live Client 与接入工厂"],
        ["apps/api/src/modules/sync/", "Inbox、管道、轮询、告警"],
        ["apps/api/src/modules/*/", "店铺/配置/目录/上架/销量/财务等"],
        ["apps/api/src/data/repo/", "仓储抽象（json / postgres 占位）"],
        [".env.example / GO_LIVE.md", "环境变量模板与上线清单"],
    ],
)

add_heading_cn("3.4 关键设计点", 2)
add_bullet("Inbox 先行：原始事件先持久化，处理失败可重试，避免「收了又丢」。")
add_bullet("幂等键：订单用 order_nr，库存用 SKU+仓库，Catalog 用 partner_sku。")
add_bullet("版本覆盖：仅当远端/导入更新时间更新时才覆盖本地。")
add_bullet("告警可见：导入失败、会话拉取失败、解析失败都会进健康中心。")
add_bullet("配置与财务解耦：财务用配置费率估算；后续可升级为下单时刻快照。")
add_bullet("采集与业务解耦：网页改版主要改 portal 解析/URL，不动 Catalog/财务页面。")

add_heading_cn("四、实现方案说明", 1)
add_heading_cn("4.1 技术选型", 2)
add_table(
    ["层级", "技术", "原因"],
    [
        ["前端", "React 19 + Vite + React Router", "模块多、多标签切换快"],
        ["后端", "Node.js + Express + TypeScript + Zod", "模块清晰，易部署"],
        [
            "平台对接",
            "NoonClient + portal/mock/live 多实现",
            "无官方 API 也能用；有 API 可插拔",
        ],
        ["门户解析", "CSV/JSON 解析 + 表头别名", "兼容后台导出差异"],
        ["本地存储", "JSON 文件（默认）", "零依赖快速起步"],
        ["目标存储", "PostgreSQL（接口已留）", "上线后换驱动，业务路由尽量不动"],
        ["包管理", "pnpm workspace", "前后端一体安装与并行启动"],
    ],
)

add_heading_cn("4.2 NOON 接入实现（门户为主）", 2)
add_para(
    "接入契约（NoonClient）统一约定：校验就绪状态、获取订单、商品、库存。"
    "portal 实现优先读「门户缓冲」（导入或会话拉取写入的数据），再进入 Inbox 管道。",
    first_line=True,
)
add_bullet("Portal 导入：POST /api/noon/portal/import（前端可上传文件或粘贴文本）")
add_bullet("Portal 会话拉取：POST /api/noon/portal/fetch（依赖 Cookie + 数据 URL）")
add_bullet("Portal 缓冲查询：GET /api/noon/portal/cache")
add_bullet("模板下载：GET /api/noon/portal/template/{orders|catalog|stock}")
add_bullet("浏览器自动化：browser-collector 占位，需要时再接 Playwright")
add_bullet("mock / live：开发演示或未来官方 API 备用通道仍保留")

add_heading_cn("4.3 业务模块实现思路", 2)
add_bullet("Catalog：读本地 catalog + stock；可配合接入页导入后再刷新。")
add_bullet("链接上架：本地额度扣减 + 任务队列。")
add_bullet("销量统计：对本地 orders 聚合 GMV、件数、状态。")
add_bullet("Best Sell：用本地订单与 Catalog 粗排爆款。")
add_bullet("财务：收入=订单×回购汇率；成本粗估空运与测评；输出利润与快照。")

add_heading_cn("4.4 防丢防错防静默失败", 2)
add_table(
    ["风险", "对策", "界面表现"],
    [
        ["重复导入同一文件", "eventId 去重，业务 UPSERT", "Inbox 标记 duplicate / 不重复记账"],
        ["CSV 表头不匹配", "别名映射；解析 0 行则报错", "导入失败红灯 + 明确文案"],
        ["Cookie 过期", "会话拉取失败告警", "红灯：需重新登录复制 Cookie"],
        ["后台页面改版", "改 URL/解析器，不动业务页", "健康中心持续报错直到修好"],
        ["处理到一半中断", "事件已在 Inbox，可继续处理", "积压数上升，黄灯"],
    ],
)

add_heading_cn("五、运行环境与启动方式", 1)
add_heading_cn("5.1 开发机要求", 2)
add_table(
    ["项目", "要求"],
    [
        ["操作系统", "Windows 10/11 或常见 Linux 服务器"],
        ["Node.js", "≥ 20"],
        ["包管理器", "pnpm（推荐）"],
        ["浏览器", "Chrome / Edge 现代版本（用于登录后台与复制 Cookie）"],
        ["项目位置", "部署在本地或服务器任意工作目录即可"],
        ["可选", "PostgreSQL（切换 DATA_DRIVER=postgres 时）"],
    ],
)

add_heading_cn("5.2 本地启动", 2)
add_para("在项目根目录终端执行：", bold=True)
add_para("pnpm install")
add_para("pnpm dev")
add_para("启动后访问：")
add_bullet("前端：http://localhost:5173")
add_bullet("后端：http://localhost:8787")
add_bullet("健康检查：GET /api/health ；就绪信息：GET /api/ready")
add_bullet("门户导入：前端「NOON 接入」页，或 POST /api/noon/portal/import")

add_heading_cn("5.3 环境变量（上线关键）", 2)
add_para("模板文件：项目根目录下的 .env.example。常用项如下：", first_line=True)
add_table(
    ["变量", "含义", "示例"],
    [
        ["NOON_MODE", "portal / mock / live / inherit", "portal"],
        ["NOON_ENABLED", "是否启用同步", "true"],
        ["NOON_COUNTRY_CODE", "站点", "sa / ae / eg"],
        ["NOON_BASE_URL", "后台相关根地址（可选）", "卖家后台域名"],
        ["DATA_DRIVER", "存储驱动", "json（默认）或 postgres"],
        ["PORT", "API 端口", "8787"],
        ["NOON_API_*", "仅 live 模式需要", "有官方 API 时再填"],
    ],
)

add_heading_cn("5.4 生产部署建议（简版）", 2)
add_bullet("前端 build 后由 Nginx 静态托管，/api 反代到 Node 服务。")
add_bullet("后端用 pm2 或 systemd 守护；开启 HTTPS。")
add_bullet("门户 Cookie 属于敏感信息，仅保存在服务端配置中，勿提交到代码仓库。")
add_bullet("日志保留 Inbox 失败事件与导入摘要，便于排查字段映射问题。")

add_heading_cn("六、项目规划与分期", 1)
add_heading_cn("6.1 分期路线", 2)
add_table(
    ["阶段", "目标", "交付物"],
    [
        [
            "P0（已完成）",
            "能跑、能配、能采集同步、能告警",
            "主页、配置、NOON 接入（portal）、同步健康、Inbox",
        ],
        [
            "P1（主体已完成）",
            "用已同步数据做日常经营",
            "Catalog、上架额度、销量、Best Sell",
        ],
        [
            "P2（进行中/占位）",
            "利润与供应链",
            "财务深化、采购补货发货、广告、测评",
        ],
        ["P3（规划）", "增长与多渠道", "选品助手、亚马逊/1688 通道"],
    ],
)

add_heading_cn("6.2 上线/日常使用最小动作清单（门户模式）", 2)
add_bullet("设置 NOON_MODE=portal（或界面选择 portal）并启动服务")
add_bullet("打开「NOON 接入」，按数据类型下载模板，对照后台导出字段")
add_bullet("从 NOON 卖家后台导出订单/商品/库存文件并导入")
add_bullet("（可选）配置 Cookie + 数据 URL，使用「会话拉取并同步」")
add_bullet("在同步健康中心确认 Inbox 处理成功、无红灯")
add_bullet("检查 Catalog、销量、财务页是否出现对应数据")

add_heading_cn("6.3 后续建议优先级", 2)
add_bullet("高：按真实后台导出样例固化字段映射；定时提醒导入/拉取")
add_bullet("高：PostgreSQL 落地，替换 JSON 文件")
add_bullet("中：需要全自动时再启用 Playwright 浏览器采集")
add_bullet("中：财务改为订单级成本快照；测评余额流水")
add_bullet("低：选品助手、广告深度、多平台上架")
add_bullet("低：若未来官方开放 API，再启用 live 通道")

add_heading_cn("七、验收标准（对照需求）", 1)
add_table(
    ["验收项", "怎么验", "通过标准"],
    [
        ["不丢数", "导入已知行数的导出文件后对比本地条数", "业务表条数对齐；失败可重导"],
        ["不错数", "重复导入同一文件；乱序旧数据", "不重复记账；旧包不覆盖新状态"],
        ["错误可见", "上传错误格式文件或过期 Cookie", "健康中心/接入页出现明确警告"],
        ["配置可用", "修改运费与测评配置并保存", "刷新后仍在；财务估算随之变化"],
        ["门户可用", "仅用导出导入（不依赖官方 API）", "业务页能显示同步后的数据"],
        ["界面可读", "非开发人员打开主页与同步页", "能判断是否正常、额度、有无告警"],
    ],
)

add_heading_cn("八、总结", 1)
add_para(
    "慧运营 的定位很明确：它是 NOON 卖家的本地经营驾驶舱。"
    "平台后台是经营真相的来源；在没有官方开放接口时，慧运营通过门户采集"
    "（导出导入 / 会话拉取）把真相稳定搬进来，再叠加上商家自己的成本规则与经营动作。",
    first_line=True,
)
add_para(
    "架构上刻意把「接触 NOON」收口到接入层，把「不能丢、不能错、出错必须看见」"
    "做成同步底座的默认能力。因此：网页可以变、导出格式可以调，"
    "但业务模块与健康中心不必推倒重来——换的是采集适配器，不是整套 ERP。",
    first_line=True,
)

add_para("相关文档（均在项目根目录）：", bold=True)
add_bullet("项目说明：README.md")
add_bullet("上线对接清单：GO_LIVE.md")
add_bullet("环境变量模板：.env.example")

footer = doc.add_paragraph()
footer.paragraph_format.space_before = Pt(18)
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("—— 文档结束 ——")
set_run_font(r, size=10, color=(120, 120, 120))

doc.save(out_path)
print(out_path)
print("size", os.path.getsize(out_path))
