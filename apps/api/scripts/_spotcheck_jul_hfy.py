# -*- coding: utf-8 -*-
"""Spot-check HFY 2026-07 against source files + pure-profit formula."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook

OUT = Path(r"D:\Desktop\gp_reports_26_06_09")
PORTAL = Path(
    r"D:\outpost-data\shops\PRJ482387\portal-imports\sales_export_2026-07-02_2026-09-29.csv"
)
FEES = Path(r"D:\Outpost_ERP\apps\api\scripts\_out\gp-sources\PRJ482387\fees.csv")
REPORT = OUT / "HFY_report_26_07.xlsx"
SPOT = Path(r"D:\Outpost_ERP\apps\api\scripts\_out\gp-sources\spotcheck_jul_hfy_verify.json")

FX = 1.79
FEE_COLS = [
    "Referral Fee including VAT",
    "Fullfilment & Logistics Fees including VAT",
    "Shipping Credits including VAT",
    "Other Order Fees including VAT",
]


def num(x) -> float:
    try:
        return float(str(x or "0").replace(",", ""))
    except ValueError:
        return 0.0


def main() -> None:
    # 1) portal revenue / units Jul
    rev = 0.0
    units = 0
    with PORTAL.open(encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            ts = r.get("order_timestamp") or ""
            if not ts.startswith("2026-07"):
                continue
            units += 1
            st = (r.get("status") or "").strip()
            if st in ("Delivered", "Shipped", "Processing"):
                rev += num(r.get("gmv_lcy") or r.get("offer_price"))

    # 2) platform fees Jul
    platform = 0.0
    with FEES.open(encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            if not (r.get("Transaction Date") or "").startswith("2026-07"):
                continue
            if (r.get("Transaction Type") or "") not in ("order", "order_update"):
                continue
            fee_sum = sum(num(r.get(c)) for c in FEE_COLS)
            platform += -fee_sum if fee_sum < 0 else fee_sum

    # 3) workbook totals
    wb = load_workbook(REPORT, data_only=True)
    ws = wb["report"]
    # find labels in col R
    book = {}
    for row in range(2, 20):
        label = ws.cell(row, 18).value
        val = ws.cell(row, 19).value
        if label:
            book[str(label)] = val

    purchase = float(book.get("采购成本(CNY)") or 0)
    intl = float(book.get("国际运费成本(CNY)") or 0)
    dom = float(book.get("国内运费成本(CNY)") or 0)
    pkg = float(book.get("产品包装采购(CNY)") or 0)
    book_rev = float(book.get("营业额(SAR)") or 0)
    book_fee = float(book.get("平台费(SAR)") or 0)
    book_pure = float(book.get("纯利(CNY)") or 0)

    net = book_rev - book_fee
    expected_pure = round(net * FX - purchase - intl - dom - pkg, 2)

    # 4) recompute purchase from SKU rows
    sku_purchase = 0.0
    sku_intl = 0.0
    sku_dom = 0.0
    sku_rev = 0.0
    for row in range(2, ws.max_row + 1):
        if not ws.cell(row, 2).value:
            continue
        sku_rev += num(ws.cell(row, 4).value)
        sku_purchase += num(ws.cell(row, 13).value)
        sku_intl += num(ws.cell(row, 14).value)
        sku_dom += num(ws.cell(row, 15).value)

    result = {
        "portal_jul_revenue": round(rev, 2),
        "portal_jul_units": units,
        "fees_jul_platform": round(platform, 2),
        "book": book,
        "sku_sum_revenue": round(sku_rev, 2),
        "sku_sum_purchase": round(sku_purchase, 2),
        "sku_sum_intl": round(sku_intl, 2),
        "sku_sum_dom": round(sku_dom, 2),
        "checks": {
            "revenue_matches_portal": abs(book_rev - rev) < 0.05,
            "revenue_matches_sku_sum": abs(book_rev - sku_rev) < 0.05,
            "platform_matches_fees_csv": abs(book_fee - platform) < 0.05,
            "purchase_matches_sku_sum": abs(purchase - sku_purchase) < 0.05,
            "intl_matches_sku_sum": abs(intl - sku_intl) < 0.05,
            "dom_matches_sku_sum": abs(dom - sku_dom) < 0.05,
            "pure_formula_ok": abs(book_pure - expected_pure) < 0.05,
        },
        "expected_pure": expected_pure,
        "delta_pure": round(book_pure - expected_pure, 4),
        "template_formula": "纯利 = (营业额-平台费)*1.79 - 采购 - 国际运费 - 国内运费 - 包装; VAT缺源未扣",
        "notes": [
            "portal Jul 无 CIR/Cancelled 等，COGS 件数=销售件数",
            "平台费来自 Transaction View，非 Statement VAT 拆分",
            "头程为装箱实付按历史箱量加权单位成本×当月销量",
        ],
    }
    SPOT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result["checks"], ensure_ascii=False, indent=2))
    print("expected_pure", expected_pure, "book_pure", book_pure)
    all_ok = all(result["checks"].values())
    print("ALL_OK" if all_ok else "HAS_FAIL")
    raise SystemExit(0 if all_ok else 2)


if __name__ == "__main__":
    main()
