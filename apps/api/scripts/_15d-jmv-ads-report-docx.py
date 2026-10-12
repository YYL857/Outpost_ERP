# -*- coding: utf-8 -*-
"""近15日 总/广告/自然 JMV与销量 → Word (.docx) 到真实桌面 D:\\Desktop"""
from __future__ import annotations

import json
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Cm

SHOPS = [
    {"id": "PRJ396770", "name": "HUI Beauty"},
    {"id": "PRJ482387", "name": "HFY-KSA"},
]
SHANGHAI = timezone(timedelta(hours=8))
DESKTOP = Path(r"D:\Desktop")
if not DESKTOP.is_dir():
    DESKTOP = Path.home() / "Desktop"


def shanghai_today() -> date:
    return datetime.now(SHANGHAI).date()


def ymd(d: date) -> str:
    return d.isoformat()


def parse_ymd(s: str) -> date:
    return date.fromisoformat(s)


def round2(n: float) -> float:
    return round(float(n or 0) + 1e-12, 2)


def pct(curr: float, prev: float):
    if prev == 0:
        return None
    return round2((curr - prev) / abs(prev) * 100)


def share(part: float, total: float):
    if not total:
        return None
    return round2(part / total * 100)


def fmt(n) -> str:
    v = float(n or 0)
    if abs(v - round(v)) < 1e-9:
        return f"{int(round(v)):,}"
    return f"{v:,.2f}"


def fmt_pct(p) -> str:
    if p is None:
        return "—"
    return f"{'+' if p > 0 else ''}{p:.1f}%"


def is_deal(status: str) -> bool:
    s = (status or "").strip().lower()
    if not s:
        return True
    import re

    return not re.search(r"cancel|void|fail|reject|refunded|returned", s)


def order_ymd(placed_at: str):
    """Parse Noon placedAt → Shanghai calendar day."""
    s = (placed_at or "").strip()
    if not s:
        return None
    try:
        # e.g. "2026-09-29 16:14:01 UTC"
        if s.endswith(" UTC"):
            dt = datetime.strptime(s, "%Y-%m-%d %H:%M:%S UTC").replace(
                tzinfo=timezone.utc
            )
            return dt.astimezone(SHANGHAI).date()
        # ISO with Z / offset
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(SHANGHAI).date()
    except Exception:
        return None


def set_run_font(run, size=10, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Microsoft YaHei"
    r = run._element.get_or_add_rPr()
    rFonts = r.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        r.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, *, bold=False, color=None, size=9, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    if align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)


ADS = RGBColor(0x92, 0x40, 0x0E)
ORG = RGBColor(0x16, 0x65, 0x34)


def build_shop(shop_id: str, name: str, days: list[str]):
    state_path = Path(rf"D:\outpost-data\shops\{shop_id}\json\state.json")
    st = json.loads(state_path.read_text(encoding="utf-8"))
    day_set = set(days)

    sales = {d: {"units": 0, "jmv": 0.0} for d in days}
    for f in st.get("salesFacts") or []:
        d = f.get("date")
        if d not in sales:
            continue
        sales[d]["units"] += f.get("units") or 0
        sales[d]["jmv"] += f.get("amount") or 0

    deals = {d: set() for d in days}
    for o in st.get("orders") or []:
        if not is_deal(o.get("status") or ""):
            continue
        d = order_ymd(o.get("placedAt") or "")
        if d is None:
            continue
        key = ymd(d)
        if key not in deals:
            continue
        nr = o.get("orderNr")
        if nr:
            deals[key].add(nr)

    ads = {
        d: {"spend": 0.0, "revenue": 0.0, "orders": 0, "clicks": 0, "facts": 0}
        for d in days
    }
    multi_skip = 0
    start, end = days[0], days[-1]
    for f in st.get("adsFacts") or []:
        fr, to = f.get("fromDate"), f.get("toDate")
        if fr != to:
            if to and fr and to >= start and fr <= end:
                multi_skip += 1
            continue
        if fr not in ads:
            continue
        ads[fr]["spend"] += f.get("spend") or 0
        ads[fr]["revenue"] += f.get("revenue") or 0
        ads[fr]["orders"] += f.get("orders") or 0
        ads[fr]["clicks"] += f.get("clicks") or 0
        ads[fr]["facts"] += 1

    series = []
    for d in days:
        jmv = round2(sales[d]["jmv"])
        units = sales[d]["units"]
        deal_orders = len(deals[d])
        ads_jmv = round2(ads[d]["revenue"])
        ads_orders = int(round(ads[d]["orders"]))
        spend = round2(ads[d]["spend"])
        org_jmv = round2(max(0, jmv - ads_jmv))
        org_orders = max(0, deal_orders - ads_orders)
        series.append(
            {
                "date": d,
                "jmv": jmv,
                "adsJmv": ads_jmv,
                "organicJmv": org_jmv,
                "units": units,
                "dealOrders": deal_orders,
                "adsOrders": ads_orders,
                "organicOrders": org_orders,
                "spend": spend,
                "clicks": ads[d]["clicks"],
                "adsFacts": ads[d]["facts"],
                "roas": round2(ads_jmv / spend) if spend > 0 else None,
                "adsJmvShare": share(ads_jmv, jmv),
                "adsOrderShare": share(ads_orders, deal_orders),
            }
        )

    tot = defaultdict(float)
    for r in series:
        for k in (
            "jmv",
            "adsJmv",
            "organicJmv",
            "units",
            "dealOrders",
            "adsOrders",
            "organicOrders",
            "spend",
            "clicks",
        ):
            tot[k] += r[k]
        tot["daysWithAds"] += 1 if r["adsFacts"] > 0 else 0
        tot["daysWithSales"] += 1 if (r["units"] > 0 or r["jmv"] > 0) else 0
    for k in ("jmv", "adsJmv", "organicJmv", "spend"):
        tot[k] = round2(tot[k])
    tot["roas"] = round2(tot["adsJmv"] / tot["spend"]) if tot["spend"] > 0 else None
    tot["acos"] = (
        round2(tot["spend"] / tot["adsJmv"] * 100) if tot["adsJmv"] > 0 else None
    )
    tot["adsJmvShare"] = share(tot["adsJmv"], tot["jmv"])
    tot["orgJmvShare"] = share(tot["organicJmv"], tot["jmv"])
    tot["adsOrderShare"] = share(tot["adsOrders"], tot["dealOrders"])
    tot["orgOrderShare"] = share(tot["organicOrders"], tot["dealOrders"])

    first7, last7 = series[:7], series[8:]

    def ssum(arr, key):
        return round2(sum(r[key] for r in arr))

    trend = {
        "jmv": pct(ssum(last7, "jmv"), ssum(first7, "jmv")),
        "adsJmv": pct(ssum(last7, "adsJmv"), ssum(first7, "adsJmv")),
        "organicJmv": pct(ssum(last7, "organicJmv"), ssum(first7, "organicJmv")),
        "units": pct(ssum(last7, "units"), ssum(first7, "units")),
        "adsOrders": pct(ssum(last7, "adsOrders"), ssum(first7, "adsOrders")),
        "organicOrders": pct(
            ssum(last7, "organicOrders"), ssum(first7, "organicOrders")
        ),
        "spend": pct(ssum(last7, "spend"), ssum(first7, "spend")),
    }

    # top SKUs
    sku = defaultdict(lambda: {"units": 0, "jmv": 0.0})
    for f in st.get("salesFacts") or []:
        d = f.get("date")
        if d not in day_set:
            continue
        psku = f.get("partnerSku") or ""
        sku[psku]["units"] += f.get("units") or 0
        sku[psku]["jmv"] += f.get("amount") or 0
    top_skus = sorted(
        [{"partnerSku": k, **v, "jmv": round2(v["jmv"])} for k, v in sku.items()],
        key=lambda x: -x["jmv"],
    )[:8]

    asku = defaultdict(lambda: {"spend": 0.0, "revenue": 0.0, "orders": 0})
    for f in st.get("adsFacts") or []:
        if f.get("fromDate") != f.get("toDate"):
            continue
        d = f.get("fromDate")
        if d not in day_set:
            continue
        psku = f.get("partnerSku") or ""
        asku[psku]["spend"] += f.get("spend") or 0
        asku[psku]["revenue"] += f.get("revenue") or 0
        asku[psku]["orders"] += f.get("orders") or 0
    top_ads = sorted(
        [
            {
                "partnerSku": k,
                "spend": round2(v["spend"]),
                "revenue": round2(v["revenue"]),
                "orders": v["orders"],
                "roas": round2(v["revenue"] / v["spend"]) if v["spend"] > 0 else None,
            }
            for k, v in asku.items()
        ],
        key=lambda x: -x["spend"],
    )[:8]

    return {
        "shop": {"id": shop_id, "name": name},
        "series": series,
        "tot": dict(tot),
        "trend": trend,
        "topSkus": top_skus,
        "topAds": top_ads,
        "multiDaySkipped": multi_skip,
    }


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13, bold=True)
    if level == 1:
        p.paragraph_format.space_after = Pt(4)
    else:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
    return p


def add_para(doc, text, *, size=9, bold=False, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_table(doc, headers, rows, col_colors=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=8, align="center")
        shade_cell(table.rows[0].cells[i], "F5F5F5")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            color = None
            if col_colors and ci in col_colors:
                color = col_colors[ci]
            align = "right" if ci > 0 else "left"
            bold = ri == len(rows) - 1 and str(row[0]) == "合计"
            set_cell_text(
                table.rows[ri + 1].cells[ci],
                val,
                size=8,
                color=color,
                align=align,
                bold=bold,
            )
    doc.add_paragraph()
    return table


def summary_block(doc, title, tot, trend):
    add_heading(doc, title, level=2)
    add_para(
        doc,
        f"总 JMV {fmt(tot['jmv'])} SAR · 件数 {fmt(tot['units'])} · 成交单 {fmt(tot['dealOrders'])} · 后7 vs 前7 {fmt_pct(trend['jmv'])}",
        size=10,
        bold=True,
    )
    add_para(
        doc,
        f"广告 JMV {fmt(tot['adsJmv'])}（占 {tot.get('adsJmvShare') or '—'}%）· 广告单 {fmt(tot['adsOrders'])}（占成交 {tot.get('adsOrderShare') or '—'}%）· 花费 {fmt(tot['spend'])} · ROAS {tot.get('roas') or '—'} · 后7 {fmt_pct(trend['adsJmv'])}",
        size=9,
        color=ADS,
    )
    add_para(
        doc,
        f"自然 JMV {fmt(tot['organicJmv'])}（占 {tot.get('orgJmvShare') or '—'}%）· 自然单 {fmt(tot['organicOrders'])}（占成交 {tot.get('orgOrderShare') or '—'}%）· 后7 自然JMV {fmt_pct(trend['organicJmv'])} · 后7 自然单 {fmt_pct(trend['organicOrders'])}",
        size=9,
        color=ORG,
    )


def series_rows(series, tot):
    rows = []
    for s in series:
        rows.append(
            [
                s["date"][5:],
                fmt(s["jmv"]),
                fmt(s["adsJmv"]),
                fmt(s["organicJmv"]),
                fmt(s["units"]),
                fmt(s["dealOrders"]),
                fmt(s["adsOrders"]),
                fmt(s["organicOrders"]),
                fmt(s["spend"]),
                "—" if s["roas"] is None else str(s["roas"]),
            ]
        )
    rows.append(
        [
            "合计",
            fmt(tot["jmv"]),
            fmt(tot["adsJmv"]),
            fmt(tot["organicJmv"]),
            fmt(tot["units"]),
            fmt(tot["dealOrders"]),
            fmt(tot["adsOrders"]),
            fmt(tot["organicOrders"]),
            fmt(tot["spend"]),
            "—" if tot.get("roas") is None else str(tot["roas"]),
        ]
    )
    return rows


HEADERS = [
    "日期",
    "总JMV",
    "广告JMV",
    "自然JMV",
    "件数",
    "成交单",
    "广告单",
    "自然单",
    "花费",
    "ROAS",
]
COL_COLORS = {2: ADS, 3: ORG, 6: ADS, 7: ORG}


def main():
    today = shanghai_today()
    end = today
    start = end - timedelta(days=14)
    days = [ymd(start + timedelta(days=i)) for i in range(15)]

    shops = [build_shop(s["id"], s["name"], days) for s in SHOPS]

    # combined
    combined = []
    for i, d in enumerate(days):
        rows = [r["series"][i] for r in shops]
        jmv = round2(sum(r["jmv"] for r in rows))
        ads_jmv = round2(sum(r["adsJmv"] for r in rows))
        org_jmv = round2(sum(r["organicJmv"] for r in rows))
        units = sum(r["units"] for r in rows)
        deal = sum(r["dealOrders"] for r in rows)
        ads_o = sum(r["adsOrders"] for r in rows)
        org_o = sum(r["organicOrders"] for r in rows)
        spend = round2(sum(r["spend"] for r in rows))
        combined.append(
            {
                "date": d,
                "jmv": jmv,
                "adsJmv": ads_jmv,
                "organicJmv": org_jmv,
                "units": units,
                "dealOrders": deal,
                "adsOrders": ads_o,
                "organicOrders": org_o,
                "spend": spend,
                "roas": round2(ads_jmv / spend) if spend > 0 else None,
            }
        )

    c_tot = defaultdict(float)
    for r in combined:
        for k in (
            "jmv",
            "adsJmv",
            "organicJmv",
            "units",
            "dealOrders",
            "adsOrders",
            "organicOrders",
            "spend",
        ):
            c_tot[k] += r[k]
    for k in ("jmv", "adsJmv", "organicJmv", "spend"):
        c_tot[k] = round2(c_tot[k])
    c_tot["roas"] = (
        round2(c_tot["adsJmv"] / c_tot["spend"]) if c_tot["spend"] > 0 else None
    )
    c_tot["adsJmvShare"] = share(c_tot["adsJmv"], c_tot["jmv"])
    c_tot["orgJmvShare"] = share(c_tot["organicJmv"], c_tot["jmv"])
    c_tot["adsOrderShare"] = share(c_tot["adsOrders"], c_tot["dealOrders"])
    c_tot["orgOrderShare"] = share(c_tot["organicOrders"], c_tot["dealOrders"])

    first7, last7 = combined[:7], combined[8:]

    def ssum(arr, key):
        return round2(sum(r[key] for r in arr))

    c_trend = {
        "jmv": pct(ssum(last7, "jmv"), ssum(first7, "jmv")),
        "adsJmv": pct(ssum(last7, "adsJmv"), ssum(first7, "adsJmv")),
        "organicJmv": pct(ssum(last7, "organicJmv"), ssum(first7, "organicJmv")),
        "units": pct(ssum(last7, "units"), ssum(first7, "units")),
        "adsOrders": pct(ssum(last7, "adsOrders"), ssum(first7, "adsOrders")),
        "organicOrders": pct(
            ssum(last7, "organicOrders"), ssum(first7, "organicOrders")
        ),
        "spend": pct(ssum(last7, "spend"), ssum(first7, "spend")),
    }

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    add_heading(doc, "近 15 日 · 总 / 广告 / 自然 · JMV 与销量", level=1)
    add_para(
        doc,
        f"窗口 {days[0]} → {days[-1]}（上海日历） · 币种 SAR · 生成于 {datetime.now(timezone.utc).isoformat()} · 店铺：HUI Beauty + HFY-KSA",
        size=9,
        color=RGBColor(0x55, 0x55, 0x55),
    )
    add_para(
        doc,
        "颜色：橙色=广告归因；绿色=自然（总 − 广告）。无单日广告导出的日期广告列为 0，自然会被抬高。",
        size=8,
        color=RGBColor(0x66, 0x66, 0x66),
    )

    summary_block(doc, "两店合计", c_tot, c_trend)
    add_para(doc, "两店合计日明细", size=10, bold=True)
    add_table(doc, HEADERS, series_rows(combined, c_tot), COL_COLORS)

    for r in shops:
        summary_block(
            doc,
            f"{r['shop']['name']}（{r['shop']['id']}）",
            r["tot"],
            r["trend"],
        )
        add_para(doc, "日明细", size=10, bold=True)
        add_table(doc, HEADERS, series_rows(r["series"], r["tot"]), COL_COLORS)

        add_para(doc, "总 JMV Top SKU", size=9, bold=True)
        add_table(
            doc,
            ["#", "PSKU", "JMV", "件"],
            [
                [str(i + 1), s["partnerSku"], fmt(s["jmv"]), fmt(s["units"])]
                for i, s in enumerate(r["topSkus"])
            ],
        )
        add_para(doc, "广告归因 Top SKU（按花费）", size=9, bold=True)
        add_table(
            doc,
            ["#", "PSKU", "花费", "归因JMV", "单", "ROAS"],
            [
                [
                    str(i + 1),
                    s["partnerSku"],
                    fmt(s["spend"]),
                    fmt(s["revenue"]),
                    fmt(s["orders"]),
                    "—" if s["roas"] is None else str(s["roas"]),
                ]
                for i, s in enumerate(r["topAds"])
            ],
            {3: ADS},
        )
        add_para(
            doc,
            f"有销量日 {int(r['tot']['daysWithSales'])}/15 · 有单日广告日 {int(r['tot']['daysWithAds'])}/15"
            + (
                f" · 跳过多日区间广告 {r['multiDaySkipped']} 条"
                if r["multiDaySkipped"]
                else ""
            ),
            size=8,
            color=RGBColor(0x66, 0x66, 0x66),
        )

    add_heading(doc, "口径说明", level=2)
    for line in [
        "总 JMV = salesFacts.amount；广告 JMV = adsFacts.revenue（仅 fromDate=toDate 单日）；自然 JMV = max(0, 总 − 广告)",
        "成交单 = orders 按下单日去重（排除取消类）；广告单 = adsFacts.orders；自然单 = max(0, 成交单 − 广告单)",
        "件数 = salesFacts.units（广告侧无件数字段，故不做件数三分）",
        "后7 vs 前7：窗口内第 9–15 日合计相对第 1–7 日（跳过第 8 日）",
        r"数据源：D:\outpost-data\shops\{PRJ}\json\state.json",
    ]:
        add_para(doc, "• " + line, size=8)

    DESKTOP.mkdir(parents=True, exist_ok=True)
    stamp = ymd(today).replace("-", "")
    ts = datetime.now(SHANGHAI).strftime("%H%M%S")
    out = DESKTOP / f"Outpost_近15日_总广告自然_JMV销量_{stamp}_{ts}.docx"
    try:
        doc.save(str(out))
    except PermissionError:
        out = DESKTOP / f"Outpost_近15日_总广告自然_JMV销量_{stamp}_{ts}_alt.docx"
        doc.save(str(out))

    # also dump json beside it
    payload = {
        "start": days[0],
        "end": days[-1],
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "combined": {"series": combined, "tot": dict(c_tot), "trend": c_trend},
        "shops": [
            {
                "shop": r["shop"],
                "tot": r["tot"],
                "trend": r["trend"],
                "series": r["series"],
                "topSkus": r["topSkus"],
                "topAds": r["topAds"],
            }
            for r in shops
        ],
    }
    json_out = DESKTOP / f"Outpost_近15日_总广告自然_JMV销量_{stamp}_{ts}.json"
    json_out.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({"docx": str(out), "json": str(json_out), "combined": dict(c_tot)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
